import io
import os

from .text_preprocessor import clean_text

MAX_BYTES = int(os.getenv("MAX_UPLOAD_MB", "5")) * 1024 * 1024
ALLOWED = {".pdf", ".docx"}


def validate_file(filename: str, data: bytes) -> str:
    """Validate extension, size and magic bytes. Returns the extension or raises ValueError."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in ALLOWED:
        raise ValueError("Only PDF and DOCX files are accepted")
    if not data:
        raise ValueError("Empty file")
    if len(data) > MAX_BYTES:
        raise ValueError(f"File exceeds {MAX_BYTES // (1024 * 1024)} MB limit")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise ValueError("File content is not a valid PDF")
    if ext == ".docx" and not data.startswith(b"PK\x03\x04"):
        raise ValueError("File content is not a valid DOCX")
    return ext


def _pdf_text(data: bytes) -> str:
    import fitz  # PyMuPDF
    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        for page in doc:
            txt = page.get_text()
            if not txt.strip() and os.getenv("ENABLE_OCR") == "1":
                txt = _ocr_page(page)
            parts.append(txt)
    return "\n".join(parts)


def _ocr_page(page) -> str:
    try:
        import pytesseract
        from PIL import Image
        pix = page.get_pixmap(dpi=200)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return pytesseract.image_to_string(img)
    except Exception:
        return ""


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, filename: str) -> str:
    ext = validate_file(filename, data)
    try:
        raw = _pdf_text(data) if ext == ".pdf" else _docx_text(data)
    except ValueError:
        raise
    except Exception as exc:  # malformed files are handled safely
        raise ValueError("Could not parse file: it may be corrupt") from exc
    return clean_text(raw)
